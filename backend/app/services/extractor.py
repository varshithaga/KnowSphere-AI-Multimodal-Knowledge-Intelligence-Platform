from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation


# ---------- PDF ----------
def extract_pdf(path: str) -> list[tuple[int | None, str]]:
    pages = []
    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            pages.append((page_number, page.get_text()))
    return pages


# ---------- Word (.docx) ----------
def extract_docx(path: str) -> list[tuple[int | None, str]]:
    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    return [(None, text)]


# ---------- PowerPoint (.pptx) ----------
def extract_pptx(path: str) -> list[tuple[int | None, str]]:
    slides = []
    presentation = Presentation(path)
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
        slides.append((slide_number, "\n".join(texts)))
    return slides


# ---------- Plain text (.txt) ----------
def extract_txt(path: str) -> list[tuple[int | None, str]]:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    return [(None, text)]


# ---------- Pick the right extractor ----------
EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "txt": extract_txt,
}


def extract_text(path: str, file_type: str) -> list[tuple[int | None, str]]:
    """Return a list of (page_number, text). page_number is None if unknown."""
    extractor = EXTRACTORS.get(file_type)
    if extractor is None:
        raise ValueError(f"No extractor for file type '{file_type}'")
    return extractor(path)